import os
import sys
from docx import Document
from fpdf import FPDF
from dotenv import load_dotenv
from .gemini import geminiApi
import PyPDF2
from pdfminer.high_level import extract_text as pdfminer_extract_text

load_dotenv(override=True)

# Función para imprimir directamente a stdout (siempre visible)
def log_info(message):
    """Imprime mensaje directamente a stdout"""
    print(message, flush=True)
    sys.stdout.flush()

def log_error(message):
    """Imprime error directamente a stdout"""
    print(message, flush=True)
    sys.stdout.flush()

class DocumentExtractor:
    def __init__(self):
        self.geminiApi = geminiApi()
        
    def convert_word_to_pdf(self, docx_path, pdf_path):
        """Convierte un archivo Word a PDF"""
        document = Document(docx_path)
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        pdf.set_font("Arial", size=12)

        for paragraph in document.paragraphs:
            # Codificar el texto para evitar problemas con caracteres especiales
            text = paragraph.text.encode('latin-1', 'replace').decode('latin-1')
            pdf.multi_cell(0, 10, text)
        
        pdf.output(pdf_path)

    def extract_text_from_pdf(self, pdf_path):
        """Extrae texto de un PDF usando PyPDF2 y pdfminer como respaldo"""
        try:
            log_info("\n" + "="*80)
            log_info(f"📄 EXTRAYENDO TEXTO DEL PDF")
            log_info("="*80)
            log_info(f"📁 Archivo: {pdf_path}")
            log_info(f"📏 Tamaño: {os.path.getsize(pdf_path)} bytes")
            
            # Intento 1: Usar PyPDF2
            text = ""
            try:
                with open(pdf_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    num_pages = len(pdf_reader.pages)
                    log_info(f"📖 Número de páginas: {num_pages}")
                    
                    for page_num in range(num_pages):
                        page = pdf_reader.pages[page_num]
                        page_text = page.extract_text()
                        text += page_text + "\n"
                        log_info(f"✅ Página {page_num + 1}/{num_pages} extraída ({len(page_text)} caracteres)")
                
                log_info(f"\n✅ Extracción con PyPDF2 completada: {len(text)} caracteres totales")
            except Exception as e:
                log_error(f"⚠️ PyPDF2 falló: {str(e)}")
            
            # Intento 2: Si PyPDF2 no extrajo suficiente texto, usar pdfminer
            if len(text.strip()) < 50:
                log_info(f"⚠️ Texto insuficiente ({len(text.strip())} caracteres)")
                log_info("🔄 Intentando con pdfminer.six...")
                try:
                    text = pdfminer_extract_text(pdf_path)
                    log_info(f"✅ Texto extraído con pdfminer.six: {len(text)} caracteres")
                except Exception as e:
                    log_error(f"❌ pdfminer.six también falló: {str(e)}")
            
            # Mostrar preview del texto extraído
            if text:
                log_info("\n" + "="*80)
                log_info("📋 PREVIEW DEL TEXTO EXTRAÍDO (primeros 800 caracteres):")
                log_info("="*80)
                preview = text[:800] + "..." if len(text) > 800 else text
                log_info(preview)
                log_info("="*80)
                log_info(f"📊 RESUMEN: {len(text)} caracteres, {len(text.split())} palabras")
                log_info("="*80 + "\n")
            else:
                log_error("❌ No se pudo extraer texto del PDF\n")
            
            self.geminiApi.configure_2()
            return text.strip()
            
        except Exception as e:
            log_error(f"❌ Error al extraer texto del PDF: {str(e)}")
            return None
        
    def extract_text_from_docx(self, docx_path):
        """Extrae texto de un archivo Word"""
        try:
            log_info("\n" + "="*80)
            log_info(f"📄 EXTRAYENDO TEXTO DEL DOCX")
            log_info("="*80)
            log_info(f"📁 Archivo: {docx_path}")
            log_info(f"📏 Tamaño: {os.path.getsize(docx_path)} bytes")
            
            document = Document(docx_path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            text = "\n".join(paragraphs)
            
            log_info(f"✅ Texto extraído: {len(text)} caracteres")
            log_info(f"📖 Número de párrafos: {len(paragraphs)}")
            
            # Mostrar preview del texto
            if text:
                log_info("\n" + "="*80)
                log_info("📋 PREVIEW DEL TEXTO EXTRAÍDO (primeros 800 caracteres):")
                log_info("="*80)
                preview = text[:800] + "..." if len(text) > 800 else text
                log_info(preview)
                log_info("="*80)
                log_info(f"📊 RESUMEN: {len(text)} caracteres, {len(text.split())} palabras")
                log_info("="*80 + "\n")
            
            self.geminiApi.configure_2()
            return text
            
        except Exception as e:
            log_error(f"❌ Error al extraer texto del DOCX: {str(e)}")
            return None
        
    def extract_text_from_txt(self, txt_path):
        """Extrae texto de un archivo TXT"""
        try:
            log_info("\n" + "="*80)
            log_info(f"📄 LEYENDO ARCHIVO TXT")
            log_info("="*80)
            log_info(f"📁 Archivo: {txt_path}")
            log_info(f"📏 Tamaño: {os.path.getsize(txt_path)} bytes")
            
            with open(txt_path, 'r', encoding='utf-8') as file:
                text = file.read()
            
            log_info(f"✅ Texto leído: {len(text)} caracteres")
            
            # Mostrar preview del texto
            if text:
                log_info("\n" + "="*80)
                log_info("📋 PREVIEW DEL TEXTO EXTRAÍDO (primeros 800 caracteres):")
                log_info("="*80)
                preview = text[:800] + "..." if len(text) > 800 else text
                log_info(preview)
                log_info("="*80)
                log_info(f"📊 RESUMEN: {len(text)} caracteres, {len(text.split())} palabras")
                log_info("="*80 + "\n")
            
            self.geminiApi.configure_2()
            return text
            
        except Exception as e:
            log_error(f"❌ Error al leer archivo TXT: {str(e)}")
            return None
        
    def extract_prioritized_requirements(self, file_path, description):
        """Extrae y prioriza requisitos de un documento"""
        log_info("\n" + "🔍 "*40)
        log_info(f"🎯 PROCESO DE PRIORIZACIÓN INICIADO")
        log_info(f"📁 Archivo: {file_path}")
        log_info(f"💬 Descripción: {description}")
        log_info("🔍 "*40 + "\n")
        
        file_extension = os.path.splitext(file_path)[1].lower()
        log_info(f"📎 Extensión detectada: {file_extension}")
        
        if file_extension == '.pdf':
            text = self.extract_text_from_pdf(file_path)
        elif file_extension == '.docx':
            text = self.extract_text_from_docx(file_path)
        elif file_extension == '.txt':
            text = self.extract_text_from_txt(file_path)
        else:
            log_error(f"❌ Formato de archivo no soportado: {file_extension}")
            return None

        if not text:
            log_error("❌ No se pudo extraer texto del archivo")
            return None

        log_info("\n" + "🤖 "*40)
        log_info("🤖 ENVIANDO A GEMINI PARA PRIORIZACIÓN...")
        log_info(f"📝 Texto a enviar: {len(text)} caracteres")
        log_info("🤖 "*40 + "\n")
        
        full_text = f"Descripción del usuario: {description}\n\n{text}"
        prioritized_requirements = self.geminiApi.generate_classification_prioritization(full_text)
        
        log_info(f"✅ Priorización completada: {len(prioritized_requirements)} requisitos")
        return prioritized_requirements
    
    def clasification_requirements(self, file_path, description):
        """Clasifica requisitos de un documento"""
        log_info("\n" + "🔍 "*40)
        log_info(f"📊 PROCESO DE CLASIFICACIÓN INICIADO")
        log_info(f"📁 Archivo: {file_path}")
        log_info(f"💬 Descripción: {description}")
        log_info("🔍 "*40 + "\n")
        
        file_extension = os.path.splitext(file_path)[1].lower()
        log_info(f"📎 Extensión detectada: {file_extension}")

        
        if file_extension == '.pdf':
            text = self.extract_text_from_pdf(file_path)
        elif file_extension == '.docx':
            text = self.extract_text_from_docx(file_path)
        elif file_extension == '.txt':
            text = self.extract_text_from_txt(file_path)
        else:
            log_error(f"❌ Formato de archivo no soportado: {file_extension}")
            return None

        if not text:
            log_error("❌ No se pudo extraer texto del archivo")
            return None

        log_info("\n" + "🤖 "*40)
        log_info("🤖 ENVIANDO A GEMINI PARA CLASIFICACIÓN...")
        log_info(f"📝 Texto a enviar: {len(text)} caracteres")
        log_info("🤖 "*40 + "\n")
        
        full_text = f"Descripción del usuario: {description}\n\n{text}"
        clasification_requirements = self.geminiApi.generate_requirements(full_text)
        
        log_info(f"✅ Clasificación completada: {len(clasification_requirements)} requisitos")
        return clasification_requirements