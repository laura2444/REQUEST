from google import genai
from PIL import Image
import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from io import BytesIO
from dotenv import load_dotenv
import PyPDF2
from pdfminer.high_level import extract_text as pdfminer_extract_text
from google.genai import types

load_dotenv(override=True)

class geminiApi:
    MAX_QUERIES = 4  # Número máximo de consultas global
    _query_count_global = 0  # Contador global de consultas

    def __init__(self):
        self.api_key = os.getenv("GEMINI_API_KEY")

        # Modelos configurables desde el archivo .env
        self.text_model_name = os.getenv(
            "GEMINI_TEXT_MODEL",
            "gemini-3.6-flash"
        )

        self.image_model_name = os.getenv(
            "GEMINI_IMAGE_MODEL",
            "gemini-3.1-flash-image"
        )

        if not self.api_key:
            raise ValueError(
                "No se encontró GEMINI_API_KEY en el archivo .env"
            )

        # Cliente oficial actual de Google GenAI
        self.client = genai.Client(api_key=self.api_key)

        # Se mantiene para conservar la estructura del proyecto
        self.model = None

    @classmethod
    def _check_limit(cls):
        if cls._query_count_global >= cls.MAX_QUERIES:
            raise Exception(
                "⚠️ Has alcanzado el límite de consultas a Gemini para esta sesión."
            )

        cls._query_count_global += 1

    def configure(self):
        """Configura el modelo Gemini para visión/imágenes."""
        self.model = self.image_model_name
        return self.model

    def configure_2(self):
        """Configura el modelo Gemini para texto."""
        self.model = self.text_model_name
        return self.model

    def _generate_text(self, prompt):
        """
        Realiza una consulta de texto a Gemini usando
        el modelo configurado en el archivo .env.
        """
        if not self.model:
            self.configure_2()

        response = self.client.models.generate_content(
            model=self.model,
            contents=prompt
        )

        return response.text

    def generate_user_story_from_text(self, user_text):
        geminiApi._check_limit()

        prompt = f"""
        A partir de la siguiente descripción:

        {user_text}

        Genera historias de usuario.

        IMPORTANTE:
        - Genera entre 5 y 8 historias de usuario.
        - Cada historia debe ocupar una sola línea.
        - NO uses Markdown.
        - NO uses títulos.
        - NO uses encabezados.
        - NO uses asteriscos.
        - NO uses guiones.
        - NO uses numeración.
        - NO uses símbolos como ###, **, -, 1., etc.
        - No agregues explicaciones antes ni después.
        - Usa exactamente este formato:

        Como [rol], quiero [acción], para [beneficio].

        Ejemplo:

        Como usuario final, quiero consultar mis productos, para conocer sus características.

        Como administrador, quiero gestionar los productos, para mantener actualizada la información.

        Descripción del producto:
        {user_text}
        """

        response_text = self._generate_text(prompt)

        historias_usuario = [
            line.strip()
            for line in response_text.split("\n")
            if line.strip()
        ]

        return historias_usuario


    def generate_requirements(self, user_text):
        geminiApi._check_limit()

        prompt = (
            "Estas son las historias de usuario:\n"
            f"{user_text}\n"
            "Conviértelas en 5 requisitos funcionales y 5 no funcionales, "
            "cada requisito en una línea separada. "
            "Empieza cada línea con 'Funcional:' o 'No funcional:'."
        )

        response_text = self._generate_text(prompt)

        requisitos = [
            line.strip()
            for line in response_text.split("\n")
            if line.strip()
        ]

        return requisitos

    def generate_classification_prioritization(self, text_R):
        geminiApi._check_limit()

        estructura = (
            "Estos son los requisitos a priorizar:\n"
            f"{text_R}\n"
            "Por favor, devuélvelos enumerados del 1 al N, y prioriza con MoSCoW (Must have, Should have, Could have, Won't have). "
            "del más importante, a importancia media, al menos importante "
            "cada requisito en una línea separada."
        )

        response_text = self._generate_text(estructura)

        prioritized_requirements = [
            req.strip()
            for req in response_text.split("\n")
            if req.strip()
        ]

        return prioritized_requirements


def generate_word_document(historias_usuario):
    """Genera un documento Word con las historias de usuario."""

    doc = Document()
    doc.add_heading("Historias de Usuario", 0)

    for historia in historias_usuario:
        doc.add_paragraph(historia)

    byte_io = BytesIO()
    doc.save(byte_io)
    byte_io.seek(0)

    return byte_io.getvalue()


def generate_pdf_document(historias_usuario):
    """Genera un documento PDF con las historias de usuario."""

    byte_io = BytesIO()

    c = canvas.Canvas(byte_io, pagesize=letter)

    width, height = letter

    c.drawString(
        100,
        height - 40,
        "Historias de Usuario"
    )

    y = height - 60

    for historia in historias_usuario:

        c.drawString(
            100,
            y,
            historia
        )

        y -= 20

        if y < 40:
            c.showPage()
            y = height - 40

    c.save()

    byte_io.seek(0)

    return byte_io.getvalue()


def generate_txt_document(historias_usuario):
    """Genera un documento TXT con las historias de usuario."""

    content = "\n".join(historias_usuario)

    return content.encode("utf-8")
