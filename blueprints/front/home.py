from flask import Blueprint, render_template, current_app, flash, send_file, request, redirect, url_for, session
from form import MyFormDocument
from werkzeug.utils import secure_filename
import os
from io import BytesIO
from classe.document import DocumentExtractor
from classe.gemini import geminiApi

import logging
logging.basicConfig(level=logging.DEBUG)

home_bp = Blueprint("home", __name__, template_folder="templates")

def print_separator(title=""):
    """Imprime un separador visual en la consola"""
    print("\n" + "="*80)
    if title:
        print(f"🔷 {title}")
        print("="*80)

def print_file_info(file_path, description):
    """Imprime información del archivo subido"""
    print_separator("INFORMACIÓN DEL ARCHIVO SUBIDO")
    print(f"📁 Ruta del archivo: {file_path}")
    print(f"📝 Nombre del archivo: {os.path.basename(file_path)}")
    print(f"📏 Tamaño del archivo: {os.path.getsize(file_path)} bytes")
    print(f"📄 Extensión: {os.path.splitext(file_path)[1]}")
    print(f"💬 Descripción del usuario: {description}")
    print("="*80 + "\n")

@home_bp.route('/')
def home():
    return render_template('home.html')

@home_bp.route('/clasificacion', methods=['GET', 'POST'])
def clasificacion():
    form = MyFormDocument()
    if form.validate_on_submit():
        print_separator("📤 NUEVA SOLICITUD DE CLASIFICACIÓN")
        
        description = form.description.data
        file = form.file.data
        
        if not description:
            flash('Por favor, proporcione una descripción.', 'error')
            print("❌ Error: No se proporcionó descripción")
        elif not file:
            flash('Por favor, suba un documento.', 'error')
            print("❌ Error: No se subió ningún archivo")
        else:
            filename = secure_filename(file.filename)
            file_path = os.path.join(current_app.config['UPLOADED_DOCUMENTS_DEST'], filename)
            
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            file.save(file_path)
            
            print(f"✅ Archivo guardado en: {file_path}")
            print_file_info(file_path, description)
            
            print("🚀 Iniciando proceso de clasificación...")
            
            try:
                extractor = DocumentExtractor()
                clasifica_requirements = extractor.clasification_requirements(file_path, description)
            except Exception as e:
                print(f"❌ Error de Gemini: {e}")
                flash('El servicio de IA no respondió. Intenta de nuevo en unos segundos.', 'error')
                return render_template('clasificacion.html', form=form)

            if clasifica_requirements:
                print_separator("✅ CLASIFICACIÓN COMPLETADA CON ÉXITO")
                print(f"📊 Total de requisitos clasificados: {len(clasifica_requirements)}")
                print("\n📋 REQUISITOS CLASIFICADOS:")
                print("-"*80)
                for i, req in enumerate(clasifica_requirements, 1):
                    print(f"{i}. {req}")
                print("-"*80 + "\n")
                
                flash('Requisitos clasificados correctamente.', 'success')
                session['textos'] = clasifica_requirements
                return render_template('clasificacion.html', form=form, textos=clasifica_requirements)
            else:
                print("❌ Error al procesar el documento")
                flash('Error al procesar el documento.', 'error')
    
    return render_template('clasificacion.html', form=form)


@home_bp.route('/priorizacion', methods=['GET', 'POST'])
def priorizacion():
    form = MyFormDocument()
    if form.validate_on_submit():
        print_separator("📤 NUEVA SOLICITUD DE PRIORIZACIÓN")
        
        description = form.description.data
        file = form.file.data
        
        if not description:
            flash('Por favor, proporcione una descripción.', 'error')
            print("❌ Error: No se proporcionó descripción")
        elif not file:
            flash('Por favor, suba un documento.', 'error')
            print("❌ Error: No se subió ningún archivo")
        else:
            filename = secure_filename(file.filename)
            file_path = os.path.join(current_app.config['UPLOADED_DOCUMENTS_DEST'], filename)
            
            os.makedirs(os.path.dirname(file_path), exist_ok=True)
            file.save(file_path)
            
            print(f"✅ Archivo guardado en: {file_path}")
            print_file_info(file_path, description)
            
            print("🚀 Iniciando proceso de priorización...")
            extractor = DocumentExtractor()
            prioritized_requirements = extractor.extract_prioritized_requirements(file_path, description)
            
            if prioritized_requirements:
                print_separator("✅ PRIORIZACIÓN COMPLETADA CON ÉXITO")
                print(f"📊 Total de requisitos priorizados: {len(prioritized_requirements)}")
                print("\n📋 REQUISITOS PRIORIZADOS (del más al menos importante):")
                print("-"*80)
                for i, req in enumerate(prioritized_requirements, 1):
                    print(f"{i}. {req}")
                print("-"*80 + "\n")
                
                flash('Requisitos priorizados obtenidos correctamente.', 'success')
                session['textos'] = prioritized_requirements
                return render_template('priorizacion.html', form=form, textos=prioritized_requirements)
            else:
                print("❌ Error al procesar el documento")
                flash('Error al procesar el documento.', 'error')

    return render_template('priorizacion.html', form=form)


@home_bp.route('/HU_imagen', methods=['GET', 'POST'])
def HU_imagen():
    form = MyFormDocument()
    if form.validate_on_submit():
        description = form.description.data.strip()

        if not description:
            flash('Por favor, ingrese una descripción.', 'error')
        else:
            api = geminiApi()
            textos = api.generate_user_story_from_text(description)

            session['textos'] = textos

            return render_template(
                'HU_imagen.html',
                form=form,
                textos=textos,
            )

    return render_template('HU_imagen.html', form=form)


@home_bp.route('/download_document/<format>', methods=['GET'])
def download_document(format):
    print_separator(f"📥 DESCARGA DE DOCUMENTO EN FORMATO: {format.upper()}")
    
    textos = session.get('textos', [])
    
    if not textos:
        print("❌ No hay contenido disponible para descargar")
        flash('No hay historias de usuario disponibles para descargar.', 'error')
        return redirect(url_for('home.priorizacion'))

    print(f"📊 Elementos a exportar: {len(textos)}")

    if format == 'word':
        document_content = generate_word_document(textos)
        filename = 'textos.docx'
        mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format == 'pdf':
        document_content = generate_pdf_document(textos)
        filename = 'textos.pdf'
        mimetype = 'application/pdf'
    elif format == 'txt':
        document_content = generate_txt_document(textos)
        filename = 'textos.txt'
        mimetype = 'text/plain'
    else:
        print(f"❌ Formato no admitido: {format}")
        return 'Formato no admitido', 400

    output_folder = current_app.config['GENERATED_UPLOADS_FOLDER']
    os.makedirs(output_folder, exist_ok=True)

    file_path = os.path.join(output_folder, filename)
    with open(file_path, 'wb') as f:
        f.write(document_content)

    print(f"✅ Documento generado: {file_path}")
    print(f"📏 Tamaño del archivo: {os.path.getsize(file_path)} bytes")
    print("="*80 + "\n")

    return send_file(file_path, as_attachment=True, download_name=filename, mimetype=mimetype)




def generate_word_document(textos):
    print("📝 Generando documento Word...")
    from docx import Document
    document = Document()
    for historia in textos:
        document.add_paragraph(historia)
    byte_io = BytesIO()
    document.save(byte_io)
    byte_io.seek(0)
    print("✅ Documento Word generado")
    return byte_io.getvalue()

def generate_pdf_document(textos):
    print("📄 Generando documento PDF...")
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    byte_io = BytesIO()
    c = canvas.Canvas(byte_io, pagesize=letter)
    y = 750
    for historia in textos:
        c.drawString(40, y, historia)
        y -= 20
    c.showPage()
    c.save()
    byte_io.seek(0)
    print("✅ Documento PDF generado")
    return byte_io.getvalue()

def generate_txt_document(textos):
    print("📋 Generando documento TXT...")
    content = "\n".join(textos)
    print("✅ Documento TXT generado")
    return content.encode('utf-8')